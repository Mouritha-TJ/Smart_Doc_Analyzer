import io
from datetime import datetime
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

def generate_pdf_report(analysis_data: dict) -> bytes:
    """
    Generates a professional PDF report using ReportLab.
    Returns binary PDF bytes.
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=40,
        leftMargin=40,
        topMargin=40,
        bottomMargin=40
    )

    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=22,
        leading=26,
        textColor=colors.HexColor('#1e1b4b'),
        alignment=0,
        spaceAfter=6
    )

    subtitle_style = ParagraphStyle(
        'DocSubTitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        textColor=colors.HexColor('#64748b'),
        spaceAfter=15
    )

    h2_style = ParagraphStyle(
        'SectionH2',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=14,
        leading=18,
        textColor=colors.HexColor('#312e81'),
        spaceBefore=14,
        spaceAfter=8
    )

    body_style = ParagraphStyle(
        'BodyDark',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#1e293b'),
        spaceAfter=6
    )

    story = []

    # Title & Subtitle
    story.append(Paragraph("SmartDoc NLP — Analysis Report", title_style))
    story.append(Paragraph(f"Document: <b>{analysis_data.get('document_name', 'Document')}</b> | Date: {datetime.now().strftime('%B %d, %Y')}", subtitle_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#6366f1'), spaceAfter=15))

    # Overview Table
    meta_table_data = [
        [
            Paragraph("<b>Document Type:</b>", body_style),
            Paragraph(analysis_data.get("classification", {}).get("category", "General"), body_style),
            Paragraph("<b>Confidence:</b>", body_style),
            Paragraph(f"{analysis_data.get('classification', {}).get('confidence', 50)}%", body_style)
        ],
        [
            Paragraph("<b>Original Words:</b>", body_style),
            Paragraph(str(analysis_data.get("original_word_count", 0)), body_style),
            Paragraph("<b>Compression:</b>", body_style),
            Paragraph(f"{analysis_data.get('compression_percentage', 0)}%", body_style)
        ],
        [
            Paragraph("<b>Vulnerability Level:</b>", body_style),
            Paragraph(f"<b>{analysis_data.get('vulnerability', {}).get('level', 'LOW')}</b> (Score: {analysis_data.get('vulnerability', {}).get('score', 0)}/100)", body_style),
            Paragraph("<b>Overall Sentiment:</b>", body_style),
            Paragraph(analysis_data.get("sentiment", {}).get("overall", "Neutral"), body_style)
        ]
    ]

    meta_table = Table(meta_table_data, colWidths=[110, 150, 110, 150])
    meta_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#f8fafc')),
        ('PADDING', (0,0), (-1,-1), 6),
        ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor('#e2e8f0')),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
    ]))
    story.append(meta_table)
    story.append(Spacer(1, 15))

    # Executive Summary
    story.append(Paragraph("1. Executive Summary", h2_style))
    story.append(Paragraph(analysis_data.get("summary", "N/A"), body_style))
    story.append(Spacer(1, 10))

    # Keyphrases & Themes
    story.append(Paragraph("2. Keyphrases & Themes", h2_style))
    kw_items = analysis_data.get("keywords", [])
    if kw_items and isinstance(kw_items[0], dict):
        kw_str = ", ".join([f"{item['phrase']}" for item in kw_items[:10]])
    else:
        kw_str = ", ".join(kw_items[:10])

    story.append(Paragraph(f"<b>Keyphrases:</b> {kw_str}", body_style))
    themes_str = ", ".join(analysis_data.get("themes", []))
    story.append(Paragraph(f"<b>Themes:</b> {themes_str}", body_style))
    story.append(Spacer(1, 10))

    # Security & Vulnerability Analysis
    vuln = analysis_data.get("vulnerability", {})
    story.append(Paragraph("3. Security & Risk Analysis", h2_style))
    story.append(Paragraph(f"Risk Level: <b>{vuln.get('level', 'LOW')}</b> (Risk Score: {vuln.get('score', 0)}/100)", body_style))

    if vuln.get("indicators"):
        story.append(Paragraph("<b>Detected Risk Indicators:</b>", body_style))
        for ind in vuln["indicators"][:5]:
            page_cite = f" [Page {ind['page']}]" if ind.get("page") else ""
            story.append(Paragraph(f"• <b>{ind['severity']}</b> ({ind['category']}): {ind['context']}{page_cite}", body_style))
    else:
        story.append(Paragraph("No security risk indicators detected.", body_style))
    story.append(Spacer(1, 10))

    # Readability & Quality
    read = analysis_data.get("readability", {})
    story.append(Paragraph("4. Readability & Quality Indicators", h2_style))
    flesch_val = int(round(float(read.get('flesch_reading_ease', 0))))
    story.append(Paragraph(f"Flesch Reading Ease: <b>{flesch_val}/100</b> ({read.get('grade_level', 'Standard')})", body_style))
    story.append(Paragraph(f"Vocabulary Richness: {read.get('vocabulary_richness', 0)} | Sentence Complexity: {read.get('sentence_complexity', 'Moderate')}", body_style))
    story.append(Spacer(1, 10))

    # Named Entities
    story.append(Paragraph("5. Extracted Named Entities", h2_style))
    ents = analysis_data.get("entities", {}).get("flat", [])
    if ents:
        ent_table_data = [["Entity Name", "Category", "Frequency"]]
        for e in ents[:12]:
            ent_table_data.append([e.get("name", ""), e.get("type", ""), str(e.get("count", 1))])

        ent_table = Table(ent_table_data, colWidths=[240, 160, 120])
        ent_table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#4338ca')),
            ('TEXTCOLOR', (0,0), (-1,0), colors.white),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('PADDING', (0,0), (-1,-1), 5),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#cbd5e1')),
        ]))
        story.append(ent_table)
    else:
        story.append(Paragraph("No named entities detected.", body_style))

    doc.build(story)
    pdf_bytes = buffer.getvalue()
    buffer.close()
    return pdf_bytes

def generate_docx_report(analysis_data: dict) -> bytes:
    """
    Generates a professional DOCX Word report using python-docx.
    Returns binary DOCX bytes.
    """
    doc = docx.Document()

    # Document Header
    title = doc.add_heading("SmartDoc NLP — Analysis Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT

    p_sub = doc.add_paragraph()
    p_sub.add_run(f"Document: {analysis_data.get('document_name', 'Document')}\n").bold = True
    p_sub.add_run(f"Generated on: {datetime.now().strftime('%B %d, %Y - %I:%M %p')}\n")

    doc.add_heading("1. Document Overview & Classification", level=1)
    cls = analysis_data.get("classification", {})
    doc.add_paragraph(f"Category: {cls.get('category', 'General Document')} (Confidence: {cls.get('confidence', 50)}%)")
    doc.add_paragraph(f"Reasoning: {cls.get('reasoning', 'N/A')}")
    doc.add_paragraph(f"Original Word Count: {analysis_data.get('original_word_count', 0)} words | Compression Rate: {analysis_data.get('compression_percentage', 0)}%")

    doc.add_heading("2. Executive Summary", level=1)
    doc.add_paragraph(analysis_data.get("summary", "N/A"))

    doc.add_heading("3. Keyphrases & Themes", level=1)
    kw_items = analysis_data.get("keywords", [])
    if kw_items and isinstance(kw_items[0], dict):
        kw_str = ", ".join([f"{item['phrase']}" for item in kw_items[:12]])
    else:
        kw_str = ", ".join(kw_items[:12])
    doc.add_paragraph(f"Keyphrases: {kw_str}")
    doc.add_paragraph(f"Themes: {', '.join(analysis_data.get('themes', []))}")

    doc.add_heading("4. Security & Risk Analysis", level=1)
    vuln = analysis_data.get("vulnerability", {})
    doc.add_paragraph(f"Vulnerability Level: {vuln.get('level', 'LOW')} (Risk Score: {vuln.get('score', 0)}/100)")
    for ind in vuln.get("indicators", [])[:5]:
        page_str = f" [Page {ind['page']}]" if ind.get("page") else ""
        doc.add_paragraph(f"[{ind['severity']}] {ind['category']}: {ind['context']}{page_str}", style='List Bullet')

    doc.add_heading("5. Readability & Quality Metrics", level=1)
    read = analysis_data.get("readability", {})
    flesch_val = int(round(float(read.get('flesch_reading_ease', 0))))
    doc.add_paragraph(f"Flesch Reading Ease: {flesch_val}/100 ({read.get('grade_level', 'Standard')})")
    doc.add_paragraph(f"Sentence Complexity: {read.get('sentence_complexity', 'Moderate')}")

    doc.add_heading("6. Named Entities", level=1)
    ents = analysis_data.get("entities", {}).get("flat", [])
    if ents:
        table = doc.add_table(rows=1, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = "Entity Name"
        hdr_cells[1].text = "Category"
        hdr_cells[2].text = "Frequency"

        for e in ents[:15]:
            row_cells = table.add_row().cells
            row_cells[0].text = e.get("name", "")
            row_cells[1].text = e.get("type", "")
            row_cells[2].text = str(e.get("count", 1))

    buffer = io.BytesIO()
    doc.save(buffer)
    docx_bytes = buffer.getvalue()
    buffer.close()
    return docx_bytes
